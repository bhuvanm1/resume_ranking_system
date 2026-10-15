"""Text extraction for common resume file types."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Extract searchable text from a PDF, DOCX or TXT resume.

    Scanned/image-only PDFs need OCR, which this beginner-friendly project does
    not enable by default. The function raises a useful error when no text exists.
    """
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(file_bytes))
            pages = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            text = "\n".join(pages)
        elif extension == ".docx":
            document = Document(BytesIO(file_bytes))
            chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            text = "\n".join(chunks)
        elif extension == ".txt":
            text = file_bytes.decode("utf-8", errors="replace")
        else:
            raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT files.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read this file ({type(exc).__name__}). Try exporting it as a text-based PDF or DOCX.") from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text.strip()) < 30:
        raise ValueError(
            "Very little text could be extracted. If this is a scanned PDF, use a text-based PDF or DOCX; OCR is not enabled in this demo."
        )
    return text
